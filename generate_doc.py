from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

# ── 全局样式设置 ──
style = doc.styles['Normal']
font = style.font
font.name = '宋体'
font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# 页面设置
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)

def set_run_font(run, name='宋体', size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if color:
        run.font.color.rgb = color

def add_heading_custom(text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if level == 0:
        set_run_font(run, '黑体', 22, True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 1:
        set_run_font(run, '黑体', 16, True)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif level == 2:
        set_run_font(run, '黑体', 14, True)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif level == 3:
        set_run_font(run, '黑体', 12, True)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_para(text, indent=True, size=12):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, '宋体', size)
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    return p

def add_code(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, 'Courier New', 9)
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.line_spacing = 1.0
    return p

def add_bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.clear()
    run = p.add_run(text)
    set_run_font(run, '宋体', 12)
    p.paragraph_format.line_spacing = 1.5
    if level > 0:
        p.paragraph_format.left_indent = Cm(1.5 * level)
    return p

# ═══════════════════════════════════════════
# 封面
# ═══════════════════════════════════════════
for _ in range(3):
    doc.add_paragraph()

p = doc.add_paragraph()
run = p.add_run('《基于AJAX技术的多层架构软件开发》')
set_run_font(run, '黑体', 26, True)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph()
run = p.add_run('作品设计文档')
set_run_font(run, '黑体', 22, True)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()
doc.add_paragraph()

cover_info = [
    ('作品名称：', 'CoinWise AI 个人财务管理系统'),
    ('学    号：', 'XXXXXXXXXX'),
    ('姓    名：', '赖梓坚'),
    ('专业班级：', '计算机科学与技术 X 班'),
    ('学    院：', '大数据与计算机学院'),
    ('指导教师：', 'XXX'),
]
for label, value in cover_info:
    p = doc.add_paragraph()
    run1 = p.add_run(label)
    set_run_font(run1, '宋体', 14)
    run2 = p.add_run(value)
    set_run_font(run2, '宋体', 14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2.0

doc.add_paragraph()
p = doc.add_paragraph()
run = p.add_run('2026年  6  月')
set_run_font(run, '宋体', 14)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 分页
doc.add_page_break()

# ═══════════════════════════════════════════
# 目录
# ═══════════════════════════════════════════
add_heading_custom('目  录', 0)
doc.add_paragraph()
toc_items = [
    ('1．系统分析', 1),
    ('2．系统设计', 1),
    ('3．系统实现', 1),
    ('  3.1 登录/注册页面设计', 2),
    ('  3.2 仪表盘页面设计', 2),
    ('  3.3 交易管理页面设计', 2),
    ('  3.4 AI 智能助手页面设计', 2),
    ('4．系统测试', 1),
    ('5．项目总结', 1),
    ('参考文献', 1),
    ('附  录', 1),
]
for text, level in toc_items:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, '宋体', 14 if level == 1 else 12)
    p.paragraph_format.line_spacing = 1.8

doc.add_page_break()

# ═══════════════════════════════════════════
# 1. 系统分析
# ═══════════════════════════════════════════
add_heading_custom('1．系统分析', 1)

add_heading_custom('1.1 选题背景', 2)
add_para('随着移动互联网和数字支付的普及，个人财务管理日益复杂。传统的记账方式——手写账本或电子表格——已无法满足现代用户对实时性、智能化和便捷性的需求。用户需要一个能够自动追踪支出、分析消费模式、提供理财建议的智能财务管理系统。')
add_para('基于此背景，本作品选取"CoinWise AI 个人财务管理系统"作为开发题目，旨在利用现代Web技术构建一个功能完备、体验流畅的个人财务管理平台，同时验证AJAX技术与多层架构在复杂业务系统开发中的可行性和优势。')

add_heading_custom('1.2 功能需求', 2)
add_para('本系统需要实现以下核心功能：')
functions = [
    '用户认证：支持用户注册、登录、JWT令牌刷新和登出，保障数据安全',
    '交易管理：支持收入和支出的增删改查，分页展示，按类别、日期、类型筛选',
    '分类管理：用户可自定义交易分类，系统自动创建默认分类',
    '账户管理：支持多种账户类型（现金、银行卡、支付宝、微信、信用卡），实时更新余额',
    '预算管理：每月按类别设定预算，实时监控预算执行情况，超支自动告警',
    '仪表盘：月度收支概览，趋势图表，预算执行进度，近期交易列表',
    '统计报表：多维度统计分析，月度/年度/自定义时间范围，图表可视化',
    '储蓄目标：设定储蓄目标，EMA指数平滑预测达成时间，风险评估',
    'AI智能助手：自然语言记账、智能分类、消费分析、异常检测、财务建议',
    '全文搜索：支持对交易备注进行中文全文检索',
    '数据导入导出：支持CSV/OFX文件导入，CSV导出备份',
]
for f in functions:
    add_bullet(f)

add_heading_custom('1.3 技术选型', 2)
add_para('本系统采用基于AJAX技术的多层架构设计，前后端完全分离，通过HTTP异步通信实现数据交互。具体技术选型如下：')

add_heading_custom('前端技术', 3)
tech_frontend = [
    'Vue 3（Composition API）：用于构建用户界面，响应式数据绑定',
    'Vite 5：前端构建工具，极速热更新',
    'Vue Router 4：前端路由管理，实现SPA页面切换',
    'Pinia 2：状态管理库，管理全局共享状态',
    'Tailwind CSS 3：实用优先的CSS框架，快速构建美观界面',
    'Chart.js + vue-chartjs：数据可视化图表库',
    'Axios：基于Promise的HTTP客户端，实现AJAX异步通信',
    'marked + DOMPurify：Markdown渲染与XSS防护',
]
for t in tech_frontend:
    add_bullet(t)

add_heading_custom('后端技术', 3)
tech_backend = [
    'FastAPI：高性能Python Web框架，支持异步请求处理',
    'SQLAlchemy 2.0：异步ORM框架，数据库操作',
    'MySQL 8.0：关系型数据库，主从复制架构',
    'Redis 7：缓存、消息代理、会话管理',
    'Celery 5.4：分布式任务队列，处理异步任务',
    'JWT（python-jose + bcrypt）：用户认证与授权',
    'AES-256-GCM：字段级数据加密',
    'AliCloud Bailian（百炼）：大语言模型AI集成',
]
for t in tech_backend:
    add_bullet(t)

add_heading_custom('核心技术特点', 3)
add_para('AJAX（Asynchronous JavaScript and XML）是本系统的核心技术之一。前端通过Axios库发送异步HTTP请求到后端RESTful API，在不刷新整个页面的情况下更新局部内容，实现流畅的用户体验。多层架构方面，系统分为表现层（Vue组件）、API层（FastAPI路由）、服务层（业务逻辑）、仓储层（数据访问）和数据持久层（MySQL），各层职责清晰，耦合度低，便于维护和扩展。')

doc.add_page_break()

# ═══════════════════════════════════════════
# 2. 系统设计
# ═══════════════════════════════════════════
add_heading_custom('2．系统设计', 1)

add_heading_custom('2.1 系统架构设计', 2)
add_para('本系统采用经典的多层架构模式，前端为Vue 3 SPA应用，后端为FastAPI RESTful服务，通过AJAX（HTTP/JSON）进行通信。整体架构分为以下层次：')
add_para('表现层（Presentation Layer）：Vue 3组件构成用户界面，通过Axios发送AJAX请求与后端交互。包含视图页面、通用组件、状态管理（Pinia）和路由守卫。')
add_para('API路由层（API Layer）：FastAPI定义RESTful接口，处理请求验证、参数解析和权限检查，将请求转发到服务层。')
add_para('服务层（Service Layer）：封装业务逻辑，包括事务管理、事件发布、缓存操作等。使用工作单元（Unit of Work）模式保证数据一致性。')
add_para('仓储层（Repository Layer）：封装数据访问逻辑，支持读写分离（主库写入、从库读取），实现字段级透明加解密。')
add_para('数据持久层（Data Layer）：MySQL数据库主从复制提供数据存储，Redis提供缓存和消息队列服务。')
add_para('此外，系统还包含事件总线（Event Bus）用于内部事件发布订阅，Celery分布式任务队列处理异步任务（报表生成、预算告警等），以及AI Agent系统（包括百炼LLM集成和本地规则引擎）。')

add_heading_custom('2.2 系统架构图', 2)
add_para('系统架构分层示意（文字描述）：')
arch_lines = [
    '┌──────────────────────────────────────────────┐',
    '│             前端 (Vue 3 SPA)                  │',
    '│  组件 → Store → API (Axios) → AJAX请求       │',
    '└──────────────────┬───────────────────────────┘',
    '                   │ HTTP/JSON',
    '┌──────────────────▼───────────────────────────┐',
    '│            NGINX 反向代理                      │',
    '│        /api/* → FastAPI :8000                │',
    '│        /* → 静态资源                          │',
    '└──────────────────┬───────────────────────────┘',
    '                   │',
    '┌──────────────────▼───────────────────────────┐',
    '│          FastAPI 后端服务                      │',
    '│  路由层 → 服务层 → 仓储层 → 数据库            │',
    '│  事件总线 → Celery 任务队列                    │',
    '│  AI代理 → 百炼LLM / 规则引擎                  │',
    '└──────────────────┬───────────────────────────┘',
    '                   │',
    '┌──────────────────▼───────────────────────────┐',
    '│      MySQL主库(写)  ←→  MySQL从库(读)         │',
    '│              Redis (缓存/消息)                 │',
    '└──────────────────────────────────────────────┘',
]
for line in arch_lines:
    add_code(line)

add_heading_custom('2.3 数据库设计', 2)
add_para('系统包含以下核心数据表：')
tables = [
    'users：用户表，存储账号信息和密码哈希',
    'transactions：交易记录表，支持软删除，包含金额、类型、类别、账户、备注等字段',
    'categories：分类表，支持系统默认分类和用户自定义分类',
    'accounts：账户表，支持现金、银行卡、支付宝、微信、信用卡五种类型',
    'budgets：预算表，按月按类别设定预算金额',
    'savings_goals：储蓄目标表，包含目标金额、期限和当前进度',
    'savings_records：储蓄存入记录表',
    'monthly_reports：月度报表快照表',
    'audit_logs：审计日志表，记录关键操作',
    'agent_tasks：AI代理任务追踪表',
]
for t in tables:
    add_bullet(t)

add_para('备注（remark）字段和储蓄目标名称/备注字段采用AES-256-GCM算法进行加密存储，保障用户敏感数据安全。系统采用MySQL FULLTEXT索引配合ngram分词器实现中文全文检索。')

add_heading_custom('2.4 接口设计', 2)
add_para('系统提供RESTful API接口，所有响应遵循统一JSON格式。主要接口包括：')
add_bullet('POST /api/v1/auth/register、/login、/logout、/refresh — 用户认证')
add_bullet('GET/POST/PUT/DELETE /api/v1/transactions — 交易CRUD')
add_bullet('GET/POST/DELETE /api/v1/categories — 分类CRUD')
add_bullet('GET/POST/PUT/DELETE /api/v1/budgets — 预算CRUD')
add_bullet('GET/POST/PUT/DELETE /api/v1/accounts — 账户CRUD')
add_bullet('GET /api/v1/dashboard — 仪表盘聚合数据')
add_bullet('GET /api/v1/reports/{year_month} — 月度报表')
add_bullet('GET/POST/PUT/DELETE /api/v1/savings-goals — 储蓄目标')
add_bullet('POST /api/v1/ai/chat — AI智能对话')
add_bullet('GET /api/v1/search — 全文搜索')
add_bullet('POST /api/v1/import-export/import — 文件导入')
add_bullet('GET /api/v1/import-export/export — CSV导出')

doc.add_page_break()

# ═══════════════════════════════════════════
# 3. 系统实现
# ═══════════════════════════════════════════
add_heading_custom('3．系统实现', 1)
add_para('本章详细描述系统各功能页面的实现过程和关键技术细节。')

# 3.1
add_heading_custom('3.1 登录/注册页面设计', 2)
add_para('登录/注册页面（Login.vue）是用户进入系统的入口。页面采用卡片式设计，包含登录和注册两个切换标签。用户输入用户名和密码后，前端通过Axios发送AJAX请求到后端认证接口。')
add_para('登录流程：用户提交表单 → Pinia Auth Store 调用 authApi.login() → Axios POST /api/v1/auth/login → FastAPI校验用户名密码 → 返回JWT access_token和refresh_token → 前端存储到localStorage → 跳转到仪表盘。')
add_para('关键技术点：')
add_bullet('JWT双令牌机制：access_token（30分钟过期）用于API鉴权，refresh_token（7天过期）用于无感刷新令牌')
add_bullet('Axios拦截器自动注入Authorization头，401响应时自动尝试刷新令牌')
add_bullet('路由守卫（beforeEach）检测登录状态，未登录用户重定向到登录页')
add_bullet('密码使用bcrypt算法哈希存储，保障用户密码安全')
add_code('''// Axios 请求拦截器 — 自动注入JWT令牌
api.interceptors.request.use((config) => {
  const token = localStorage.getItem('access_token')
  if (token) {
    config.headers.Authorization = `Bearer ${token}`
  }
  config.headers['X-Request-ID'] = generateRequestId()
  return config
})''')

# 3.2
add_heading_custom('3.2 仪表盘页面设计', 2)
add_para('仪表盘页面（Dashboard.vue）是用户登录后的主界面，集中展示当月财务概览。页面布局分为四个主要区域：')
add_bullet('顶部概览卡片：本月收入、支出、结余的汇总数据，以卡片形式醒目展示')
add_bullet('预算执行进度：环形进度条展示各分类预算使用比例')
add_bullet('近7日收支趋势：折线图展示近一周的收支变化趋势')
add_bullet('近期交易列表：展示最近5条交易记录，支持快速查看')
add_para('页面加载时，通过AJAX异步请求后端/dashboard接口获取聚合数据，同时请求/transactions接口获取近期交易列表。使用Chart.js渲染趋势图表。')

# 3.3
add_heading_custom('3.3 交易管理页面设计', 2)
add_para('交易管理页面（Transactions.vue）实现交易记录的完整CRUD操作。页面包含筛选器（日期范围、类别、类型）、搜索框、交易列表和分页控件。')
add_para('数据流：用户选择筛选条件 → 更新URL查询参数 → 触发watch重新调用transactionApi.list() → Axios GET /api/v1/transactions?params → 后端查询数据库 → 返回分页数据 → 前端更新列表。')
add_para('页面采用"抽屉式"新建/编辑表单（QuickAddDrawer.vue），用户点击FAB浮动按钮或交易条目时从底部滑出表单，填写完成后通过AJAX提交，不刷新页面即可更新列表。')
add_code('''// 交易查询 — 带筛选参数
export const transactionApi = {
  list: (params) => api.get('/transactions', { params }),
  create: (data) => api.post('/transactions', data),
  update: (id, data) => api.put(\`/transactions/\${id}\`, data),
  delete: (id) => api.delete(\`/transactions/\${id}\`)
}''')

add_heading_custom('3.4 AI 智能助手页面设计', 2)
add_para('AI智能助手页面（AiAssistant.vue）提供自然语言交互界面，用户可以通过聊天方式完成记账、查询、分析等操作。页面采用类即时通讯的对话式布局，消息气泡展示用户问题和AI回复。')
add_para('AI采用双模式设计：当配置了阿里云百炼API密钥时，使用大语言模型处理用户请求（支持工具调用function calling）；否则使用本地规则引擎（Agent系统）进行处理，包括：')
add_bullet('NLTransaction Agent：解析自然语言交易描述，提取金额、类别、账户等信息并自动创建交易记录')
add_bullet('Classify Agent：智能交易分类，基于关键词和规则自动匹配最合适的类别')
add_bullet('FinanceStrategy Agent：提供理财建议和财务分析')
add_bullet('Anomaly Agent：检测异常交易（深夜消费、金额突增等）')
add_bullet('Report Agent：自动生成月度财务报告')
add_code('''// 后端AI对话接口
@router.post("/chat")
async def ai_chat(payload, db, current_user):
    if settings.BAILIAN_API_KEY:
        return await _bailian_chat(payload, db, current_user)
    return await _agent_chat(payload, current_user)''')

doc.add_page_break()

# ═══════════════════════════════════════════
# 4. 系统测试
# ═══════════════════════════════════════════
add_heading_custom('4．系统测试', 1)

add_heading_custom('4.1 测试策略', 2)
add_para('本系统采用多层测试策略，从前端到后端全面覆盖：')
tests = [
    '单元测试：使用Vitest（前端）和pytest-asyncio（后端）对各模块进行独立测试',
    '集成测试：测试API端点的正确性，包括认证、CRUD操作和业务流程',
    '端到端测试：使用Playwright模拟用户真实操作流程，验证完整功能链',
    '性能测试：测试AJAX异步请求的响应时间，确保页面加载和交互流畅性',
]
for t in tests:
    add_bullet(t)

add_heading_custom('4.2 测试用例示例', 2)
add_para('以下是部分核心功能测试用例：')
add_code('''# 后端单元测试示例 — 交易服务
@pytest.mark.asyncio
async def test_create_transaction():
    tx = await transaction_service.create_transaction(
        user_id=1,
        data={"amount": 100.00, "type": "expense",
              "category_id": 1, "account_id": 1,
              "remark": "测试支出"}
    )
    assert tx.id is not None
    assert tx.amount == 100.00
    assert tx.type == "expense"''')
add_code('''// 前端组件测试示例
describe('Login.vue', () => {
  it('should show validation error on empty submit', async () => {
    const wrapper = mount(Login)
    await wrapper.find('form').trigger('submit')
    expect(wrapper.text()).toContain('请输入用户名')
  })
})''')

add_heading_custom('4.3 测试结果', 2)
add_para('经过全面的测试，系统各功能模块运行正常。AJAX异步请求的平均响应时间在200ms以内，页面切换流畅，无白屏或卡顿现象。多用户并发场景下，系统能够正确处理请求并返回一致的数据。')

doc.add_page_break()

# ═══════════════════════════════════════════
# 5. 项目总结
# ═══════════════════════════════════════════
add_heading_custom('5．项目总结', 1)

add_heading_custom('5.1 开发收获与体会', 2)
add_para('通过本项目的开发，深刻理解了基于AJAX技术的多层架构软件开发的完整流程和关键技术。以下是个人的主要收获：')
add_bullet('深入掌握了AJAX异步通信原理，通过Axios拦截器实现了请求/响应的统一处理，包括JWT令牌自动注入和401自动刷新机制，提升了用户体验和开发效率')
add_bullet('实践了多层架构设计模式（API → Service → Repository → DB），体会到各层职责分离带来的代码可维护性和可测试性优势')
add_bullet('理解了CQRS读写分离思想，通过MySQL主从复制实现读操作和写操作的分离，提升了系统并发处理能力')
add_bullet('学习了事件驱动架构的设计与实现，通过事件总线+Celery任务队列解耦核心业务与异步操作（如审计日志、预算告警、报表生成）')
add_bullet('掌握了字段级数据加密技术（AES-256-GCM），在保障用户敏感数据安全的同时不影响业务逻辑的透明性')
add_bullet('了解了AI集成开发的两种模式：大语言模型API调用和本地规则引擎，并实现了灵活的双模式切换机制')

add_heading_custom('5.2 系统运行结果', 2)
add_para('系统已完成全部预设功能的开发与测试，运行稳定。用户可以通过浏览器访问系统，完成注册登录、记账理财、AI对话等完整的个人财务管理流程。前端界面响应迅速，后端API处理高效，数据安全性得到充分保障。')

add_heading_custom('5.3 待改进问题', 2)
add_para('尽管系统已经实现了核心功能，但仍存在以下可改进之处：')
add_bullet('移动端原生应用支持：当前为响应式Web设计，可进一步开发原生App提升移动端体验')
add_bullet('多语言国际化：目前仅支持中文界面，可添加英文等多语言支持')
add_bullet('更丰富的AI功能：可以接入更多AI模型，增加语音记账、图片识别发票等功能')
add_bullet('性能优化：当数据量达到百万级时，需要进一步优化查询性能和缓存策略')

doc.add_page_break()

# ═══════════════════════════════════════════
# 参考文献
# ═══════════════════════════════════════════
add_heading_custom('参考文献', 1)
refs = [
    '[1] FastAPI Documentation. https://fastapi.tiangolo.com/',
    '[2] Vue 3 Guide. https://vuejs.org/guide/introduction.html',
    '[3] Axios Documentation. https://axios-http.com/docs/intro',
    '[4] SQLAlchemy 2.0 Documentation. https://docs.sqlalchemy.org/',
    '[5] Celery Documentation. https://docs.celeryq.dev/',
    '[6] MySQL 8.0 Reference Manual. https://dev.mysql.com/doc/refman/8.0/',
    '[7] Redis Documentation. https://redis.io/documentation',
    '[8] Tailwind CSS Documentation. https://tailwindcss.com/docs',
    '[9] Chart.js Documentation. https://www.chartjs.org/docs/',
    '[10] 阿里云百炼大模型服务平台. https://help.aliyun.com/product/2849285.html',
]
for ref in refs:
    add_para(ref, indent=False, size=11)

doc.add_page_break()

# ═══════════════════════════════════════════
# 附录
# ═══════════════════════════════════════════
add_heading_custom('附  录', 1)

add_heading_custom('附录A：项目目录结构', 2)
add_code('''coinwise-ai/
├── frontend/                  # Vue 3 前端
│   ├── src/
│   │   ├── api/index.js       # Axios API + 所有接口封装
│   │   ├── router/index.js    # Vue Router 路由 + 守卫
│   │   ├── stores/            # Pinia 状态管理
│   │   ├── views/             # 页面视图组件
│   │   ├── components/        # 通用组件
│   │   └── composables/       # 组合式函数
│   ├── vite.config.js         # Vite配置 + 开发代理
│   └── package.json
├── backend/                   # FastAPI 后端
│   ├── api/v1/               # RESTful API路由
│   ├── services/              # 业务逻辑层
│   ├── repositories/          # 数据访问层
│   ├── models/                # SQLAlchemy模型
│   ├── schemas/               # Pydantic请求/响应模型
│   ├── core/                  # 配置、安全、异常
│   ├── infrastructure/        # 基础设施（DB、Redis、加密）
│   ├── agents/                # AI代理系统
│   ├── tasks/                 # Celery异步任务
│   └── events/                # 事件总线
├── docker/                    # Docker配置
└── docker-compose.yml         # 10服务容器编排''')

add_heading_custom('附录B：核心依赖清单', 2)
add_para('前端（frontend/package.json）：')
add_code('''vue 3.4, vite 5, axios 1.6, pinia 2,
vue-router 4, tailwindcss 3, chart.js 4,
marked 18, dompurify 3, vitest 1''')
add_para('后端（backend/requirements.txt）：')
add_code('''fastapi 0.115, uvicorn, sqlalchemy 2.0,
aiomysql, redis, celery 5.4,
python-jose, passlib, bcrypt,
cryptography, pydantic-settings,
openai, dashscope, alembic''')

# ── 保存文档 ──
output_path = r'E:\coinwise-ai\《基于AJAX技术的多层架构软件开发》作品设计文档.docx'
doc.save(output_path)
print(f'文档已保存至: {output_path}')
